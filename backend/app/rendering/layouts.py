"""Small helpers and renderers for the supported presentation layouts."""

import json
from collections.abc import Mapping, Sequence
from itertools import zip_longest
from pathlib import Path
from typing import Any

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

from app.models.presentation import ComponentType, Slide, SlideComponent, SlideLayout

NAVY = RGBColor(25, 43, 70)
BLUE = RGBColor(48, 105, 170)
TEAL = RGBColor(31, 142, 143)
INK = RGBColor(43, 53, 66)
MUTED = RGBColor(102, 116, 132)
PALE = RGBColor(239, 244, 249)
WHITE = RGBColor(255, 255, 255)
GOLD = RGBColor(238, 176, 70)


def component_text(component: SlideComponent) -> str:
    """Convert common JSON component shapes into readable plain text."""
    value = component.content
    if isinstance(value, str):
        return value
    if isinstance(value, Sequence) and not isinstance(value, (str, bytes)):
        return "\n".join(_value_text(item) for item in value)
    if isinstance(value, Mapping):
        if component.type == ComponentType.TITLE:
            for key in ("title", "text", "heading", "header", "name", "label", "value"):
                item = value.get(key)
                if item is not None and str(item).strip():
                    return _value_text(item)
        if component.type == ComponentType.BULLET_LIST:
            items = bullet_items(component)
            if items:
                return "\n".join(items)
        for key in ("text", "description", "summary", "content", "value", "items", "bullets", "points", "title", "label", "context"):
            item = value.get(key)
            if item is not None and str(item).strip():
                if isinstance(item, Sequence) and not isinstance(item, (str, bytes)):
                    return "\n".join(_value_text(i) for i in item if _value_text(i).strip())
                return _value_text(item)
        return "\n".join(_value_text(item) for item in value.values() if _value_text(item).strip())
    return str(value) if value is not None else ""


def _value_text(value: Any) -> str:
    if isinstance(value, Mapping):
        preferred = [value.get(key) for key in ("title", "text", "label", "value", "description")]
        return " — ".join(str(item) for item in preferred if item is not None)
    if isinstance(value, Sequence) and not isinstance(value, (str, bytes)):
        return ", ".join(_value_text(item) for item in value)
    return str(value) if value is not None else ""


def find_component(slide: Slide, *types: ComponentType) -> SlideComponent | None:
    return next((component for component in slide.components if component.type in types), None)


def title_for(slide: Slide, presentation_title: str) -> str:
    title = find_component(slide, ComponentType.TITLE)
    return component_text(title) if title else (presentation_title if slide.order == 1 else slide.purpose)


def add_text(slide, text: str, x: float, y: float, w: float, h: float,
             size: int = 18, color: RGBColor = INK, bold: bool = False,
             align: PP_ALIGN = PP_ALIGN.LEFT, font: str = "Aptos",
             margin: float = 0.05):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def add_heading(slide, title: str, page_number: int) -> None:
    add_text(slide, title, 0.72, 0.38, 11.6, 0.65, 27, NAVY, True)
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.75), Inches(1.14), Inches(0.72), Inches(0.055))
    rule.fill.solid()
    rule.fill.fore_color.rgb = TEAL
    rule.line.fill.background()
    add_text(slide, f"{page_number:02d}", 12.1, 7.08, 0.5, 0.2, 9, MUTED, align=PP_ALIGN.RIGHT)


def add_card(slide, x: float, y: float, w: float, h: float, heading: str, body: str, accent=BLUE) -> None:
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    card.fill.solid()
    card.fill.fore_color.rgb = PALE
    card.line.color.rgb = RGBColor(222, 230, 238)
    stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h))
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = accent
    stripe.line.fill.background()
    add_text(slide, heading, x + 0.24, y + 0.22, w - 0.48, 0.58, 18, NAVY, True)
    add_text(slide, body, x + 0.24, y + 0.9, w - 0.48, h - 1.05, 14, INK)


def bullet_items(component: SlideComponent | None) -> list[str]:
    if component is None:
        return []
    content = component.content
    if isinstance(content, Mapping):
        for key in ("items", "bullets", "points", "lines", "list", "text", "values", "content"):
            if key in content and content[key]:
                content = content[key]
                break
        else:
            content = []
    if isinstance(content, Sequence) and not isinstance(content, (str, bytes)):
        return [_value_text(item) for item in content if _value_text(item).strip()]
    return [line.strip(" •-\t") for line in str(content).splitlines() if line.strip(" •-\t")]


def _body_components(slide: Slide) -> list[SlideComponent]:
    return [c for c in slide.components if c.type != ComponentType.TITLE]


def _add_bullets(slide, items: list[str], x: float, y: float, w: float, h: float, size: int = 19) -> None:
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.12)
    tf.margin_top = Inches(0.08)
    for index, item in enumerate(items):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.text = f"•  {item}"
        p.level = 0
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = INK
        p.space_after = Pt(16)


def render_layout(slide, model: Slide, presentation_title: str) -> None:
    """Render one validated slide using its controlled layout."""
    layout = model.layout
    if layout == SlideLayout.TITLE:
        title = title_for(model, presentation_title)
        add_text(slide, title, 1.0, 2.05, 11.3, 1.25, 36, NAVY, True, PP_ALIGN.CENTER)
        body = find_component(model, ComponentType.TEXT)
        if body:
            add_text(slide, component_text(body), 1.7, 3.55, 9.9, 1.1, 20, MUTED, align=PP_ALIGN.CENTER)
        add_text(slide, model.purpose, 1.4, 6.75, 10.5, 0.35, 11, MUTED, align=PP_ALIGN.CENTER)
        return

    add_heading(slide, title_for(model, presentation_title), model.order)
    body = _body_components(model)
    if layout == SlideLayout.CONCLUSION:
        bullet = find_component(model, ComponentType.BULLET_LIST)
        text_components = [c for c in body if c.type in (ComponentType.TEXT, ComponentType.CARD)]
        summary = "\n\n".join(component_text(c) for c in text_components if component_text(c).strip())
        items = bullet_items(bullet) if bullet else []
        if items and summary:
            summary_box = add_text(slide, summary, 1.0, 1.55, 11.0, 1.75, 20, INK)
            summary_box.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
            _add_bullets(slide, items, 1.0, 3.45, 10.8, 2.85, 19)
        elif items:
            _add_bullets(slide, items, 1.0, 1.55, 10.8, 4.9, 21)
        elif summary:
            summary_box = add_text(slide, summary, 1.0, 1.65, 11.0, 4.8, 22, INK)
            summary_box.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        _ensure_slide_has_body_content(slide, model)
        return

    if layout == SlideLayout.TITLE_BULLETS:
        bullet = find_component(model, ComponentType.BULLET_LIST)
        items = bullet_items(bullet) if bullet else [component_text(c) for c in body if c.type in (ComponentType.TEXT, ComponentType.CARD)]
        if items:
            _add_bullets(slide, items, 1.0, 1.65, 10.8, 4.9, 19)
        _ensure_slide_has_body_content(slide, model)
        return

    if layout == SlideLayout.TWO_COLUMN:
        columns = [c for c in body if c.type in (ComponentType.TEXT, ComponentType.CARD, ComponentType.BULLET_LIST)]
        for i, col in enumerate(columns[:2]):
            x = 0.8 if i == 0 else 6.85
            w = 5.65
            heading, text = _card_content(col, i + 1)
            add_card(slide, x, 1.65, w, 4.9, heading, text, BLUE if i == 0 else TEAL)
        _ensure_slide_has_body_content(slide, model)
        return

    if layout in (SlideLayout.THREE_CARD, SlideLayout.COMPARISON):
        if layout == SlideLayout.COMPARISON:
            table = find_component(model, ComponentType.TABLE)
            if table:
                _render_table(slide, table, 0.85, 1.65, 11.65, 4.95)
                _ensure_slide_has_body_content(slide, model)
                return
        cards = [c for c in body if c.type in (ComponentType.CARD, ComponentType.TEXT, ComponentType.STATISTIC, ComponentType.BULLET_LIST)]
        cards = cards[:3]
        if not cards:
            cards = body[:3]
        width = 3.7 if len(cards) == 3 else 5.4
        gap = 0.35
        start = (13.333 - (len(cards) * width + max(0, len(cards) - 1) * gap)) / 2
        for i, item in enumerate(cards):
            heading, body_text = _card_content(item, i + 1)
            add_card(slide, start + i * (width + gap), 1.95, width, 4.25, heading, body_text, [BLUE, TEAL, GOLD][i % 3])
        _ensure_slide_has_body_content(slide, model)
        return

    if layout == SlideLayout.IMAGE_TEXT:
        image = find_component(model, ComponentType.IMAGE)
        x, y, w, h = 0.85, 1.65, 5.55, 4.95
        path = _image_path(image)
        if path and Path(path).is_file():
            try:
                slide.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w), height=Inches(h))
            except (OSError, ValueError) as exc:
                _image_placeholder(slide, x, y, w, h, f"Image unavailable\n{exc}")
        else:
            _image_placeholder(slide, x, y, w, h, "Image placeholder\nAdd a valid local image path")
        text_components = [c for c in body if c.type not in (ComponentType.IMAGE,)]
        text = "\n\n".join(component_text(c) for c in text_components)
        add_text(slide, text or model.purpose, 6.8, 1.8, 5.65, 4.5, 18, INK)
        _ensure_slide_has_body_content(slide, model)
        return

    if layout == SlideLayout.TIMELINE:
        comp = find_component(model, ComponentType.TIMELINE)
        entries = comp.content if comp else [component_text(c) for c in body]
        if not isinstance(entries, Sequence) or isinstance(entries, (str, bytes)):
            entries = [entries]
        entries = list(entries)[:5]
        if not entries:
            entries = [model.purpose]
        y = 1.85
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.35), Inches(y + 0.15), Inches(0.045), Inches(min(4.5, len(entries) * 0.88)))
        line.fill.solid(); line.fill.fore_color.rgb = TEAL; line.line.fill.background()
        for i, item in enumerate(entries):
            yy = y + i * 0.88
            dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.22), Inches(yy + 0.03), Inches(0.3), Inches(0.3))
            dot.fill.solid(); dot.fill.fore_color.rgb = BLUE; dot.line.fill.background()
            add_text(slide, _value_text(item), 1.85, yy, 10.1, 0.68, 17, INK)
        _ensure_slide_has_body_content(slide, model)
        return

    if layout == SlideLayout.STATISTIC:
        comp = find_component(model, ComponentType.STATISTIC)
        content = comp.content if comp else (body[0].content if body else {})
        data = content if isinstance(content, Mapping) else {"value": _value_text(content)}
        add_text(slide, str(data.get("value", "—")), 1.0, 1.75, 11.3, 1.45, 48, BLUE, True, PP_ALIGN.CENTER)
        label = str(data.get("label", data.get("title", "Key statistic")))
        label_size = 22 if len(label) <= 70 else 19 if len(label) <= 150 else 16
        label_box = add_text(slide, label, 1.25, 3.25, 10.8, 1.9, label_size, NAVY, True, PP_ALIGN.CENTER)
        label_box.text_frame.word_wrap = True
        label_box.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        context = data.get("context", data.get("description", ""))
        if context:
            context_box = add_text(slide, str(context), 2.0, 5.45, 9.3, 0.9, 15, MUTED, align=PP_ALIGN.CENTER)
            context_box.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        _ensure_slide_has_body_content(slide, model)
        return

    raise ValueError(f"Unsupported slide layout: {layout!r}")


def _image_path(component: SlideComponent | None) -> str | None:
    if component is None:
        return None
    content = component.content
    if isinstance(content, str):
        return content
    if isinstance(content, Mapping):
        for key in ("path", "local_path", "image_path", "src"):
            value = content.get(key)
            if isinstance(value, str):
                return value
    for key in ("path", "local_path", "image_path"):
        value = component.properties.get(key)
        if isinstance(value, str):
            return value
    return None


def _image_placeholder(slide, x: float, y: float, w: float, h: float, label: str) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = PALE
    shape.line.color.rgb = RGBColor(212, 223, 233)
    shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    run = shape.text_frame.paragraphs[0].add_run()
    run.text = label
    run.font.name = "Aptos"; run.font.size = Pt(14); run.font.color.rgb = MUTED


def _render_table(slide, component: SlideComponent, x: float, y: float, w: float, h: float) -> None:
    data = _table_content(component.content)
    raw_headers = data.get("headers", data.get("columns"))
    raw_rows = data.get("rows", [])
    headers, column_values = _table_headers_and_columns(_decode_json_value(raw_headers))
    rows = _table_rows(_decode_json_value(raw_rows), len(headers))
    if not rows and column_values:
        rows = [list(row) for row in zip_longest(*column_values, fillvalue="")]
    column_count = max([len(headers), *(len(row) for row in rows)], default=0) or 2
    if not headers:
        headers = [f"Column {index + 1}" for index in range(column_count)]
    elif len(headers) < column_count:
        headers.extend(f"Column {index + 1}" for index in range(len(headers), column_count))
    rows = [row + [""] * (column_count - len(row)) for row in rows]
    row_count = len(rows) + 1
    table = slide.shapes.add_table(row_count, column_count, Inches(x), Inches(y), Inches(w), Inches(h)).table
    for col, header in enumerate(headers):
        cell = table.cell(0, col); cell.text = str(header)
        cell.fill.solid(); cell.fill.fore_color.rgb = NAVY
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                r.font.bold = True; r.font.color.rgb = WHITE; r.font.size = Pt(13)
    for row_idx, row in enumerate(rows, start=1):
        for col in range(column_count):
            cell = table.cell(row_idx, col); cell.text = _value_text(row[col])
            cell.fill.solid(); cell.fill.fore_color.rgb = WHITE if row_idx % 2 else PALE
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.name = "Aptos"; r.font.color.rgb = INK; r.font.size = Pt(12)


def _card_content(component: SlideComponent, position: int) -> tuple[str, str]:
    """Return the provided title and text from a card, bullet-list, or text component."""
    content = component.content
    if isinstance(content, Mapping):
        nested = content.get("content")
        if isinstance(nested, Mapping):
            content = nested
        title = None
        for key in ("title", "heading", "header", "name", "label"):
            if key in content and content[key]:
                title = content[key]
                break
        heading = _value_text(title).strip() if title is not None else ""

        body_parts: list[str] = []
        # Check for direct text / description / summary / value
        for key in ("text", "description", "summary", "value"):
            if key in content and content[key]:
                val = _value_text(content[key]).strip()
                if val:
                    body_parts.append(val)
                break

        # Check for bullet / list items
        for key in ("items", "bullets", "points", "lines", "list", "values"):
            if key in content and content[key]:
                items = content[key]
                if isinstance(items, Sequence) and not isinstance(items, (str, bytes)):
                    for itm in items:
                        t = _value_text(itm).strip()
                        if t:
                            body_parts.append(f"•  {t}" if not t.startswith("•") else t)
                elif isinstance(items, str):
                    for line in items.splitlines():
                        t = line.strip(" •-\t")
                        if t:
                            body_parts.append(f"•  {t}")
                break

        body_text = "\n".join(body_parts)
        if not body_text:
            other_vals = [
                f"{k.replace('_', ' ').capitalize()}: {_value_text(v)}"
                for k, v in content.items()
                if k not in ("title", "heading", "header", "name", "label") and _value_text(v).strip()
            ]
            if other_vals:
                body_text = "\n".join(other_vals)
            elif not heading:
                body_text = component_text(component)

        return heading or f"Point {position}", body_text

    if isinstance(content, Sequence) and not isinstance(content, (str, bytes)):
        if len(content) == 1 and isinstance(content[0], Mapping):
            nested_comp = SlideComponent(
                id=component.id,
                type=component.type,
                content=content[0],
                properties=component.properties,
            )
            return _card_content(nested_comp, position)
        bullets = [_value_text(item).strip() for item in content if _value_text(item).strip()]
        body_text = "\n".join(f"•  {b}" if not b.startswith("•") else b for b in bullets)
        return f"Point {position}", body_text

    if isinstance(content, str):
        return f"Point {position}", content.strip()

    return f"Point {position}", component_text(component)


def _ensure_slide_has_body_content(slide, model: Slide) -> None:
    """Defensive validation: confirm non-title slides contain non-empty rendered body content."""
    for shape in slide.shapes:
        if shape.has_table:
            if any(cell.text.strip() for row in shape.table.rows for cell in row.cells):
                return
        elif shape.has_text_frame:
            text = shape.text_frame.text.strip()
            if shape.top is not None and shape.top > Inches(1.2) and text:
                if not (len(text) <= 3 and text.isdigit()):
                    return
    raise ValueError(f"Slide '{model.id}' layout '{model.layout.value}' rendered with empty body content.")


def _table_headers_and_columns(raw_headers: Any) -> tuple[list[str], list[list[Any]]]:
    """Normalize string headers or column descriptors with optional values."""
    if isinstance(raw_headers, Mapping):
        result = []
        for name, column_values in raw_headers.items():
            values = column_values.get("values", column_values.get("items", [])) if isinstance(column_values, Mapping) else column_values
            result.append({"name": name, "values": values})
        raw_headers = result
    if not isinstance(raw_headers, Sequence) or isinstance(raw_headers, (str, bytes)):
        return [], []
    headers: list[str] = []
    values: list[list[Any]] = []
    for index, item in enumerate(raw_headers):
        if isinstance(item, Mapping):
            headers.append(str(item.get("header", item.get("title", item.get("name", item.get("label", f"Column {index + 1}"))))))
            column_values = item.get("values", item.get("items", []))
            values.append(list(column_values) if isinstance(column_values, Sequence) and not isinstance(column_values, (str, bytes)) else [])
        else:
            headers.append(_value_text(item))
            values.append([])
    return headers, values


def _table_rows(raw_rows: Any, column_count: int) -> list[list[Any]]:
    """Normalize nested rows, mappings, or a flat row-major sequence."""
    if isinstance(raw_rows, Mapping):
        raw_rows = list(raw_rows.values())
    if not isinstance(raw_rows, Sequence) or isinstance(raw_rows, (str, bytes)):
        return [[raw_rows]] if raw_rows is not None else []
    items = list(raw_rows)
    if not items:
        return []
    if all(not isinstance(item, (Mapping, Sequence)) or isinstance(item, (str, bytes)) for item in items):
        width = max(column_count, 1)
        return [items[index:index + width] for index in range(0, len(items), width)]
    rows: list[list[Any]] = []
    for item in items:
        if isinstance(item, Mapping):
            rows.append(list(item.values()))
        elif isinstance(item, Sequence) and not isinstance(item, (str, bytes)):
            rows.append(list(item))
        else:
            rows.append([item])
    return rows


def _table_content(content: Any) -> Mapping[str, Any]:
    """Read direct, nested, or JSON-encoded table data without losing its rows."""
    value = _decode_json_value(content)
    if not isinstance(value, Mapping):
        return {}
    if any(key in value for key in ("headers", "columns", "rows")):
        return value
    for key in ("table", "table_data", "data", "content", "value"):
        nested = _decode_json_value(value.get(key))
        if isinstance(nested, Mapping) and any(name in nested for name in ("headers", "columns", "rows")):
            return nested
    raise ValueError("Table component content must contain headers or columns and rows as structured JSON.")


def _decode_json_value(value: Any) -> Any:
    """Decode structured JSON strings while leaving ordinary text untouched."""
    if isinstance(value, str):
        try:
            return json.loads(value)
        except json.JSONDecodeError:
            return value
    return value
