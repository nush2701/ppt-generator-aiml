"""Convert a validated Presentation model into a deterministic .pptx file."""

from pathlib import Path

from pptx import Presentation as PowerPoint
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt

from app.models.presentation import Presentation
from app.rendering.layouts import render_layout


class PptRenderError(RuntimeError):
    """Raised when a presentation cannot be rendered or saved."""


def render_presentation(presentation: Presentation, output_path: str | Path) -> Path:
    """Render the validated model and return the saved PowerPoint path."""
    if not isinstance(presentation, Presentation):
        raise TypeError("presentation must be a validated app.models.presentation.Presentation")
    target = Path(output_path)
    try:
        target.parent.mkdir(parents=True, exist_ok=True)
        deck = PowerPoint()
        deck.slide_width = Inches(13.333)
        deck.slide_height = Inches(7.5)
        blank = deck.slide_layouts[6]
        for item in sorted(presentation.slides, key=lambda slide: slide.order):
            slide = deck.slides.add_slide(blank)
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
            render_layout(slide, item, presentation.title)
            _add_footer(slide, presentation.title)
        deck.core_properties.title = presentation.title
        deck.core_properties.subject = f"Presentation {presentation.id}"
        deck.core_properties.author = "PPT Generator"
        deck.save(target)
    except (OSError, ValueError, KeyError, IndexError, TypeError) as exc:
        raise PptRenderError(f"Could not render presentation '{presentation.id}' to '{target}': {exc}") from exc
    return target


def _add_footer(slide, title: str) -> None:
    # A subtle footer keeps the deck visually consistent without affecting content.
    shape = slide.shapes.add_textbox(Inches(0.75), Inches(7.08), Inches(10.8), Inches(0.2))
    shape.text_frame.clear()
    shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = shape.text_frame.paragraphs[0]
    run = paragraph.add_run(); run.text = title
    run.font.name = "Aptos"; run.font.size = Pt(8); run.font.color.rgb = RGBColor(150, 160, 171)
