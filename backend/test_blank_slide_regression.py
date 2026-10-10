"""Deterministic regression tests for renderer bullet lists, tables, and layout-aware validation.

Verifies:
1. Bullet-list items render into shapes in comparison, two-column, and three-card layouts.
2. Structured table data is preserved in comparison layouts.
3. Actual exported PPTX files contain the expected bullet text, unaffected slides remain unchanged,
   and original files are preserved.
4. Validation is layout-aware: title-only slides are accepted, while slides with empty body content
   (or only heading keys) are rejected.
"""

import hashlib
from pathlib import Path
from pptx import Presentation as PowerPoint

from app.models.change import ChangeType
from app.models.presentation import (
    ComponentType,
    Presentation,
    PresentationMetadata,
    Slide,
    SlideComponent,
    SlideLayout,
)
from app.rendering import render_presentation
from app.services.slide_regenerator import SlideRegenerator


def create_initial_fixture() -> Presentation:
    """Create a controlled 5-slide presentation fixture."""
    return Presentation(
        id="fixture_pres_clean_energy",
        title="Renewable Energy Overview",
        metadata=PresentationMetadata(root={"topic": "clean energy"}),
        slides=[
            Slide(
                id="slide-01",
                order=1,
                layout=SlideLayout.TITLE,
                purpose="Introduce the topic",
                components=[
                    SlideComponent(
                        id="s1_title",
                        type=ComponentType.TITLE,
                        content={"text": "Renewable Energy Overview"},
                    ),
                    SlideComponent(
                        id="s1_sub",
                        type=ComponentType.TEXT,
                        content={"text": "Clean power for tomorrow"},
                    ),
                ],
            ),
            Slide(
                id="slide-02",
                order=2,
                layout=SlideLayout.TWO_COLUMN,
                purpose="Showcase two pillar technologies with bullet lists",
                components=[
                    SlideComponent(
                        id="s2_title",
                        type=ComponentType.TITLE,
                        content={"text": "Core Clean Technologies"},
                    ),
                    SlideComponent(
                        id="s2_col1",
                        type=ComponentType.BULLET_LIST,
                        content={
                            "title": "Solar Advantages",
                            "items": [
                                "Abundant solar resource across regions",
                                "Rapidly falling photovoltaic panel costs",
                            ],
                        },
                    ),
                    SlideComponent(
                        id="s2_col2",
                        type=ComponentType.BULLET_LIST,
                        content={
                            "title": "Wind Advantages",
                            "items": [
                                "Scalable modern turbine sizes",
                                "High overnight electricity generation",
                            ],
                        },
                    ),
                ],
            ),
            Slide(
                id="slide-03",
                order=3,
                layout=SlideLayout.COMPARISON,
                purpose="Initial comparison overview",
                components=[
                    SlideComponent(
                        id="s3_title",
                        type=ComponentType.TITLE,
                        content={"text": "Initial Energy Comparison"},
                    ),
                    SlideComponent(
                        id="s3_card1",
                        type=ComponentType.CARD,
                        content={
                            "title": "Preliminary Overview",
                            "text": "Baseline comparison before detailed analysis.",
                        },
                    ),
                ],
            ),
            Slide(
                id="slide-04",
                order=4,
                layout=SlideLayout.THREE_CARD,
                purpose="Three key resources with bullet points",
                components=[
                    SlideComponent(
                        id="s4_title",
                        type=ComponentType.TITLE,
                        content={"text": "Three Renewable Resources"},
                    ),
                    SlideComponent(
                        id="s4_card1",
                        type=ComponentType.BULLET_LIST,
                        content={
                            "title": "Solar Power",
                            "items": ["Photovoltaic arrays", "Concentrated solar power"],
                        },
                    ),
                    SlideComponent(
                        id="s4_card2",
                        type=ComponentType.CARD,
                        content={
                            "title": "Wind Power",
                            "text": "Onshore and offshore utility turbines.",
                        },
                    ),
                    SlideComponent(
                        id="s4_card3",
                        type=ComponentType.BULLET_LIST,
                        content={
                            "title": "Hydro Power",
                            "items": ["Run-of-river systems", "Pumped hydro storage"],
                        },
                    ),
                ],
            ),
            Slide(
                id="slide-05",
                order=5,
                layout=SlideLayout.COMPARISON,
                purpose="Detailed metrics table",
                components=[
                    SlideComponent(
                        id="s5_title",
                        type=ComponentType.TITLE,
                        content={"text": "Resource Comparison Table"},
                    ),
                    SlideComponent(
                        id="s5_table",
                        type=ComponentType.TABLE,
                        content={
                            "headers": ["Metric", "Solar", "Wind"],
                            "rows": [
                                ["Levelized Cost", "$36/MWh", "$38/MWh"],
                                ["Capacity Factor", "24%", "35%"],
                                ["Footprint", "Medium", "Variable"],
                            ],
                        },
                    ),
                ],
            ),
        ],
    )


def extract_slide_texts(slide) -> list[str]:
    """Collect all non-empty text strings from shapes in a slide."""
    texts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                cleaned = paragraph.text.strip()
                if cleaned:
                    texts.append(cleaned)
        elif shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    cleaned = cell.text.strip()
                    if cleaned:
                        texts.append(cleaned)
    return texts


def file_sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_renderer_and_selective_pptx_export() -> None:
    generated_dir = Path(__file__).parent / "generated"
    generated_dir.mkdir(exist_ok=True)
    v1_path = generated_dir / "regression_fixture_v1.pptx"
    v2_path = generated_dir / "regression_fixture_v2.pptx"

    try:
        # Step 1: Render v1 fixture
        initial_deck = create_initial_fixture()
        render_presentation(initial_deck, v1_path)
        assert v1_path.is_file(), "v1 PPTX was not created"
        v1_hash_initial = file_sha256(v1_path)
        v1_size_initial = v1_path.stat().st_size

        v1_pptx = PowerPoint(v1_path)
        assert len(v1_pptx.slides) == 5, f"Expected 5 slides, got {len(v1_pptx.slides)}"

        # Verify Slide 2 two-column bullet items rendered
        s2_texts = extract_slide_texts(v1_pptx.slides[1])
        assert any("Abundant solar resource across regions" in t for t in s2_texts), f"Slide 2 missing solar bullet: {s2_texts}"
        assert any("Rapidly falling photovoltaic panel costs" in t for t in s2_texts), f"Slide 2 missing solar bullet: {s2_texts}"
        assert any("Scalable modern turbine sizes" in t for t in s2_texts), f"Slide 2 missing wind bullet: {s2_texts}"
        assert any("High overnight electricity generation" in t for t in s2_texts), f"Slide 2 missing wind bullet: {s2_texts}"

        # Verify Slide 4 three-card bullet items rendered
        s4_texts = extract_slide_texts(v1_pptx.slides[3])
        assert any("Photovoltaic arrays" in t for t in s4_texts), f"Slide 4 missing card bullet: {s4_texts}"
        assert any("Concentrated solar power" in t for t in s4_texts), f"Slide 4 missing card bullet: {s4_texts}"
        assert any("Onshore and offshore utility turbines" in t for t in s4_texts), f"Slide 4 missing card text: {s4_texts}"
        assert any("Run-of-river systems" in t for t in s4_texts), f"Slide 4 missing card bullet: {s4_texts}"
        assert any("Pumped hydro storage" in t for t in s4_texts), f"Slide 4 missing card bullet: {s4_texts}"

        # Verify Slide 5 structured table rendered
        s5_texts = extract_slide_texts(v1_pptx.slides[4])
        for expected in ("Metric", "Solar", "Wind", "Levelized Cost", "$36/MWh", "$38/MWh", "Capacity Factor", "24%", "35%"):
            assert expected in s5_texts, f"Slide 5 table missing '{expected}': {s5_texts}"

        # Step 2: Simulate selective modification of Slide 3 (exact shape of Solar vs Wind Gemini output)
        modified_slide_3 = Slide(
            id="slide-03",
            order=3,
            layout=SlideLayout.COMPARISON,
            purpose="Compare solar energy and wind energy using clear points",
            components=[
                SlideComponent(
                    id="s3_title",
                    type=ComponentType.TITLE,
                    content={"text": "Solar Energy vs. Wind Energy"},
                ),
                SlideComponent(
                    id="s3_comp_solar",
                    type=ComponentType.BULLET_LIST,
                    content={
                        "title": "Solar Energy",
                        "items": [
                            "High initial installation and inverter cost",
                            "Intermittent daylight generation pattern",
                            "Modular scalability from rooftop to utility",
                        ],
                    },
                ),
                SlideComponent(
                    id="s3_comp_wind",
                    type=ComponentType.BULLET_LIST,
                    content={
                        "title": "Wind Energy",
                        "items": [
                            "Dependent on geographic wind speed patterns",
                            "Mechanical moving parts require maintenance",
                            "High capacity factor during night hours",
                        ],
                    },
                ),
            ],
        )

        modified_deck = initial_deck.model_copy(deep=True)
        modified_deck.slides[2] = modified_slide_3

        render_presentation(modified_deck, v2_path)
        assert v2_path.is_file(), "v2 PPTX was not created"

        # Step 3: Verify v1 file is preserved and completely untouched
        assert v1_path.is_file(), "Original v1 file was removed!"
        assert file_sha256(v1_path) == v1_hash_initial, "Original v1 file was modified!"
        assert v1_path.stat().st_size == v1_size_initial, "Original v1 file size changed!"

        # Step 4: Verify actual exported PPTX content in v2
        v2_pptx = PowerPoint(v2_path)
        assert len(v2_pptx.slides) == 5, f"Expected 5 slides in v2, got {len(v2_pptx.slides)}"

        # Verify Slide 3 contains ALL expected bullet points and headings
        s3_v2_texts = extract_slide_texts(v2_pptx.slides[2])
        assert any("Solar Energy vs. Wind Energy" in t for t in s3_v2_texts), f"Slide 3 missing title: {s3_v2_texts}"
        assert any("Solar Energy" in t for t in s3_v2_texts), f"Slide 3 missing Solar Energy heading: {s3_v2_texts}"
        assert any("Wind Energy" in t for t in s3_v2_texts), f"Slide 3 missing Wind Energy heading: {s3_v2_texts}"
        for expected_bullet in (
            "High initial installation and inverter cost",
            "Intermittent daylight generation pattern",
            "Modular scalability from rooftop to utility",
            "Dependent on geographic wind speed patterns",
            "Mechanical moving parts require maintenance",
            "High capacity factor during night hours",
        ):
            assert any(expected_bullet in t for t in s3_v2_texts), (
                f"Slide 3 missing expected bullet '{expected_bullet}'. Actual texts: {s3_v2_texts}"
            )

        # Step 5: Verify unaffected slides (1, 2, 4, 5) remain identical
        for index in (0, 1, 3, 4):
            v1_slide_texts = extract_slide_texts(v1_pptx.slides[index])
            v2_slide_texts = extract_slide_texts(v2_pptx.slides[index])
            assert v1_slide_texts == v2_slide_texts, (
                f"Slide {index + 1} changed between v1 and v2! v1={v1_slide_texts}, v2={v2_slide_texts}"
            )

        print("Renderer & PPTX export verification: PASSED (all bullet text verified in PPTX, unaffected slides preserved, original file preserved).")
    finally:
        v1_path.unlink(missing_ok=True)
        v2_path.unlink(missing_ok=True)


def test_layout_aware_validation() -> None:
    # 1. Valid title-only layout: MUST PASS
    title_only_slide = Slide(
        id="slide-01",
        order=1,
        layout=SlideLayout.TITLE,
        purpose="Title slide",
        components=[
            SlideComponent(
                id="c1",
                type=ComponentType.TITLE,
                content={"text": "Clean Energy Revolution"},
            ),
        ],
    )
    SlideRegenerator._validate_replacement(
        replacement=title_only_slide,
        current_slide=title_only_slide,
        allow_layout_change=False,
    )

    # 2. Non-title slide without any body components: MUST FAIL
    comp_without_body = Slide(
        id="slide-03",
        order=3,
        layout=SlideLayout.COMPARISON,
        purpose="Compare sources",
        components=[
            SlideComponent(
                id="c1",
                type=ComponentType.TITLE,
                content={"text": "Comparison"},
            ),
        ],
    )
    try:
        SlideRegenerator._validate_replacement(
            replacement=comp_without_body,
            current_slide=comp_without_body,
            allow_layout_change=False,
        )
    except ValueError as exc:
        assert "requires body components" in str(exc), f"Unexpected message: {exc}"
    else:
        raise AssertionError("Validation should have rejected comparison slide with no body components")

    # 3. Non-title slide with cards that only have headings and NO body content: MUST FAIL
    # (This is the exact bug pattern where non-empty heading dictionary misled shallow validation)
    comp_heading_only = Slide(
        id="slide-03",
        order=3,
        layout=SlideLayout.COMPARISON,
        purpose="Compare sources",
        components=[
            SlideComponent(
                id="c1",
                type=ComponentType.TITLE,
                content={"text": "Comparison"},
            ),
            SlideComponent(
                id="c2",
                type=ComponentType.CARD,
                content={"title": "Solar Energy"},
            ),
            SlideComponent(
                id="c3",
                type=ComponentType.CARD,
                content={"title": "Wind Energy"},
            ),
        ],
    )
    try:
        SlideRegenerator._validate_replacement(
            replacement=comp_heading_only,
            current_slide=comp_heading_only,
            allow_layout_change=False,
        )
    except ValueError as exc:
        assert "has no meaningful body content" in str(exc), f"Unexpected message: {exc}"
    else:
        raise AssertionError("Validation should have rejected comparison slide with only headings and no body content")

    # 4. Non-title slide with valid bullet list body content: MUST PASS
    comp_valid = Slide(
        id="slide-03",
        order=3,
        layout=SlideLayout.COMPARISON,
        purpose="Compare sources",
        components=[
            SlideComponent(
                id="c1",
                type=ComponentType.TITLE,
                content={"text": "Comparison"},
            ),
            SlideComponent(
                id="c2",
                type=ComponentType.BULLET_LIST,
                content={"title": "Solar Energy", "items": ["Modular", "Clean"]},
            ),
            SlideComponent(
                id="c3",
                type=ComponentType.BULLET_LIST,
                content={"title": "Wind Energy", "items": ["High capacity", "Scalable"]},
            ),
        ],
    )
    SlideRegenerator._validate_replacement(
        replacement=comp_valid,
        current_slide=comp_valid,
        allow_layout_change=False,
    )

    print("Layout-aware validation verification: PASSED (title-only accepted, heading-only body rejected, valid body accepted).")


def main() -> None:
    test_layout_aware_validation()
    test_renderer_and_selective_pptx_export()
    print("ALL DETERMINISTIC REGRESSION TESTS PASSED!")


if __name__ == "__main__":
    main()
