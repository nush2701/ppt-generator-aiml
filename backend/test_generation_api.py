"""API validation and real-Gemini full-regeneration route integration checks."""

from pathlib import Path

from fastapi.testclient import TestClient
from pptx import Presentation as PowerPoint

from app.main import app
from app.models.metrics import OperationType
from app.rendering import render_presentation
from app.services.metrics import MetricsStore
from app.services.presentation_store import PresentationStore
from test_full_regeneration import make_presentation


def main() -> None:
    client = TestClient(app)
    empty_prompt = client.post("/api/generate", json={"prompt": "   "})
    assert empty_prompt.status_code == 422, empty_prompt.text

    oversized_prompt = client.post("/api/generate", json={"prompt": "x" * 20_001})
    assert oversized_prompt.status_code == 422, oversized_prompt.text

    health = client.get("/health")
    assert health.status_code == 200, health.text
    assert health.json() == {"status": "ok"}

    initial_generation = client.post(
        "/api/generate",
        json={
            "prompt": (
                "Create a 5-slide professional presentation about renewable energy "
                "for college students, including an introduction, comparison, "
                "future outlook, and conclusion."
            )
        },
    )
    assert initial_generation.status_code == 200, initial_generation.text
    generated_body = initial_generation.json()
    generated_presentation, generated_version = PresentationStore().load_latest(
        generated_body["presentation_id"]
    )
    assert generated_version == 1
    assert len(generated_presentation.slides) == generated_body["slide_count"]
    generated_pptx = Path(__file__).parent / "generated" / generated_body["file_name"]
    assert generated_pptx.is_file() and len(PowerPoint(generated_pptx).slides) == generated_body["slide_count"]

    missing_presentation = client.post(
        "/api/modify",
        json={"presentation_id": "not-stored-presentation", "prompt": "Change slide 1."},
    )
    assert missing_presentation.status_code == 404, missing_presentation.text

    source = make_presentation()
    store = PresentationStore()
    initial_json = store.save(source, version=1)
    initial_bytes = initial_json.read_bytes()
    initial_pptx = initial_json.parent / f"presentation_{source.id}_v1.pptx"
    render_presentation(source, initial_pptx)
    initial_pptx_bytes = initial_pptx.read_bytes()

    ambiguous = client.post(
        "/api/modify/full",
        json={"presentation_id": source.id, "prompt": "Change it."},
    )
    assert ambiguous.status_code == 200, ambiguous.text
    ambiguous_body = ambiguous.json()
    assert ambiguous_body["status"] == "clarification_needed", ambiguous_body
    assert ambiguous_body["version"] is None
    assert ambiguous_body["file_name"] is None
    assert not (initial_json.parent / f"presentation_{source.id}_v2.json").exists()

    full = client.post(
        "/api/modify/full",
        json={
            "presentation_id": source.id,
            "prompt": "Change slide 3 to compare solar and wind energy instead.",
        },
    )
    assert full.status_code == 200, full.text
    full_body = full.json()
    assert full_body["status"] == "regenerated", full_body
    assert full_body["success"] is True
    assert full_body["presentation_id"] == source.id
    assert full_body["version"] == 2
    assert full_body["regenerated_slides"] == [f"slide-{number}" for number in range(1, 6)]
    full_pptx = initial_json.parent / full_body["file_name"]
    assert full_body["file_name"].startswith("full_regeneration_")
    assert full_pptx.is_file() and full_pptx.stat().st_size > 0
    assert len(PowerPoint(full_pptx).slides) == 5
    version_two_json = initial_json.parent / f"presentation_{source.id}_v2.json"
    assert version_two_json.is_file()
    assert initial_json.read_bytes() == initial_bytes, "Original JSON version was overwritten."
    assert initial_pptx.read_bytes() == initial_pptx_bytes, "Original PPTX version was overwritten."

    selective_ambiguous = client.post(
        "/api/modify",
        json={"presentation_id": source.id, "prompt": "Change it."},
    )
    assert selective_ambiguous.status_code == 200, selective_ambiguous.text
    selective_ambiguous_body = selective_ambiguous.json()
    assert selective_ambiguous_body["status"] == "clarification_needed"
    assert selective_ambiguous_body["regenerated_slides"] == []

    selective_success = client.post(
        "/api/modify",
        json={
            "presentation_id": source.id,
            "prompt": (
                "Rewrite slide-3, titled Renewable Energy Comparison, "
                "to compare solar and wind energy."
            ),
        },
    )
    assert selective_success.status_code == 200, selective_success.text
    selective_body = selective_success.json()
    assert selective_body["status"] == "regenerated", selective_body
    assert selective_body["regenerated_slides"] == ["slide-3"]

    records = MetricsStore().load_all()
    records_by_id = {record.run_id: record for record in records}
    run_ids = {
        generated_body["run_id"],
        ambiguous_body["run_id"],
        full_body["run_id"],
        selective_ambiguous_body["run_id"],
        selective_body["run_id"],
    }
    assert run_ids.issubset(records_by_id), "One or more API runs were not persisted to metrics.jsonl."
    assert records_by_id[generated_body["run_id"]].operation == OperationType.GENERATION
    assert records_by_id[generated_body["run_id"]].regenerated_slide_count == generated_body["slide_count"]
    assert records_by_id[ambiguous_body["run_id"]].operation == OperationType.FULL_REGENERATION
    assert records_by_id[ambiguous_body["run_id"]].status == "clarification_needed"
    assert records_by_id[ambiguous_body["run_id"]].llm_calls == 1
    assert records_by_id[full_body["run_id"]].regenerated_slide_count == 5
    assert records_by_id[full_body["run_id"]].affected_slide_ids == [f"slide-{n}" for n in range(1, 6)]
    assert records_by_id[selective_ambiguous_body["run_id"]].status == "clarification_needed"
    assert records_by_id[selective_ambiguous_body["run_id"]].regenerated_slide_count == 0
    assert records_by_id[selective_body["run_id"]].regenerated_slide_count == 1
    assert records_by_id[selective_body["run_id"]].affected_slide_ids == ["slide-3"]
    assert all(
        record.latency_ms >= 0
        and record.llm_calls >= 0
        and record.input_tokens >= 0
        and record.output_tokens >= 0
        and record.total_tokens >= 0
        and record.estimated_cost_usd >= 0
        for record in records_by_id.values()
        if record.run_id in run_ids
    )

    print(
        "API validation passed: prompt limits return HTTP 422, health returns ok, "
        "unknown IDs return HTTP 404, ambiguous full changes do not save, and "
        "generation/full/selective API runs saved versioned outputs and metrics."
    )
    for run_id in run_ids:
        metric = records_by_id[run_id]
        print(
            f"Metrics {metric.operation.value} ({metric.status}): "
            f"calls={metric.llm_calls}, tokens={metric.total_tokens}, "
            f"usage_available={metric.token_usage_available}, "
            f"unavailable_calls={metric.usage_unavailable_calls}"
        )


if __name__ == "__main__":
    main()
