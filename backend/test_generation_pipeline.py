"""Real Gemini-to-PowerPoint integration test for the initial MVP pipeline."""

from pathlib import Path
from uuid import uuid4

from pptx import Presentation as PowerPoint

from app.models.presentation import Presentation
from app.services.presentation_generator import generate_presentation


def main() -> None:
    prompt = (
        "Create a 5-slide professional presentation about renewable energy for "
        "college students. Include an introduction, major renewable sources, "
        "a comparison, a future outlook, and a conclusion."
    )
    output = Path(__file__).parent / "generated" / f"pipeline_test_{uuid4().hex}.pptx"

    result = generate_presentation(prompt, output)
    presentation = result.presentation
    assert isinstance(presentation, Presentation), "Pipeline did not return a Presentation model"
    assert len(presentation.slides) == 5, f"Expected 5 slides, received {len(presentation.slides)}"
    assert len({slide.id for slide in presentation.slides}) == 5, "Slide IDs are not unique"
    assert [slide.order for slide in presentation.slides] == [1, 2, 3, 4, 5], "Slide orders are invalid"
    assert all(slide.components for slide in presentation.slides), "At least one slide has no components"
    assert result.output_path.is_file(), f"PowerPoint was not created: {result.output_path}"
    assert result.output_path.stat().st_size > 0, "PowerPoint file is empty"

    reopened = PowerPoint(result.output_path)
    assert len(reopened.slides) == 5, f"Rendered PPT has {len(reopened.slides)} slides instead of 5"

    print(f"Presentation title: {presentation.title}")
    print(f"Slide count: {len(presentation.slides)}")
    print(f"Output path: {result.output_path}")
    print(f"File size: {result.output_path.stat().st_size:,} bytes")
    print("End-to-end Gemini generation and PowerPoint rendering completed successfully.")


if __name__ == "__main__":
    main()
