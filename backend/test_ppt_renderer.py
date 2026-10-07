"""Deterministic smoke test for all controlled PowerPoint layouts."""

import json
from importlib.metadata import version
from pathlib import Path

from pptx import Presentation as PowerPoint
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches

from app.models.presentation import (
    ComponentType,
    Presentation,
    PresentationMetadata,
    Slide,
    SlideComponent,
    SlideLayout,
)
from app.rendering import render_presentation


def component(identifier: str, kind: ComponentType, content, properties=None) -> SlideComponent:
    return SlideComponent(id=identifier, type=kind, content=content, properties=properties or {})


def sample_presentation() -> Presentation:
    layouts = [
        SlideLayout.TITLE,
        SlideLayout.TITLE_BULLETS,
        SlideLayout.TWO_COLUMN,
        SlideLayout.COMPARISON,
        SlideLayout.THREE_CARD,
        SlideLayout.IMAGE_TEXT,
        SlideLayout.TIMELINE,
        SlideLayout.STATISTIC,
        SlideLayout.CONCLUSION,
    ]
    slides = []
    for index, layout in enumerate(layouts, start=1):
        components = [component(f"slide{index}_title", ComponentType.TITLE, {"text": layout.value.replace("_", " ").title()})]
        if layout == SlideLayout.TITLE_BULLETS:
            components.append(component(f"slide{index}_bullets", ComponentType.BULLET_LIST, {"items": ["Clear result", "Useful next step"]}))
        elif layout == SlideLayout.TWO_COLUMN:
            components.extend([
                component(f"slide{index}_left", ComponentType.CARD, {"title": "Benefits", "text": "Clear and accessible."}),
                component(f"slide{index}_right", ComponentType.CARD, {"title": "Considerations", "text": "Use with care."}),
            ])
        elif layout == SlideLayout.COMPARISON:
            components.append(component(f"slide{index}_table", ComponentType.TABLE, {
                "headers": ["Source", "Primary Driver", "Availability"],
                "rows": [
                    "Solar", "Sunlight", "Daytime only",
                    "Wind", "Air currents", "Variable",
                    "Hydro", "Water flow", "Consistent",
                    "Geothermal", "Heat from the Earth", "Location-dependent",
                    "Biomass", "Organic material", "Fuel availability",
                    "Tidal", "Ocean tides", "Coastal variation",
                    "Wave", "Ocean surface motion", "Site-dependent",
                ],
            }))
        elif layout == SlideLayout.THREE_CARD:
            cards = [
                ("Solar", "Harnessing sunlight for electricity and heat."),
                ("Wind", "Using moving air to drive turbines."),
                ("Hydro", "Generating power from flowing water."),
            ]
            components.extend([component(f"slide{index}_card{card}", ComponentType.CARD, {"title": title, "text": text}) for card, (title, text) in enumerate(cards, start=1)])
        elif layout == SlideLayout.IMAGE_TEXT:
            components.extend([
                component(f"slide{index}_image", ComponentType.IMAGE, {"path": "missing_optional_image.png"}),
                component(f"slide{index}_text", ComponentType.TEXT, {"text": "Image slides can include local imagery."}),
            ])
        elif layout == SlideLayout.TIMELINE:
            components.append(component(f"slide{index}_timeline", ComponentType.TIMELINE, ["Start", "Build", "Deliver"]))
        elif layout == SlideLayout.STATISTIC:
            components.append(component(f"slide{index}_stat", ComponentType.STATISTIC, {
                "value": "42%",
                "label": "A deliberately long statistic label that checks wrapping and fitting within its assigned presentation area without extending beyond the slide or overlapping the context below it.",
                "context": "Illustrative test content.",
            }))
        elif layout == SlideLayout.CONCLUSION:
            components.extend([
                component(f"slide{index}_summary", ComponentType.TEXT, {"text": "Renewable energy can support a more diverse electricity system.\n\nProgress depends on matching each source to local conditions."}),
                component(f"slide{index}_bullets", ComponentType.BULLET_LIST, {"items": ["Use resources suited to each region", "Continue improving storage and grid planning"]}),
            ])
        elif layout == SlideLayout.TITLE:
            components.append(component(f"slide{index}_subtitle", ComponentType.TEXT, {"text": "A deterministic rendering smoke test"}))
        slides.append(Slide(id=f"slide_{index:02d}", order=index, layout=layout, purpose=f"Demonstrate {layout.value} layout", components=components))
    return Presentation(id="render_smoke_test", title="PowerPoint Rendering Smoke Test", metadata=PresentationMetadata(root={}), slides=slides)


def table_presentation(identifier: str, table_content) -> Presentation:
    """Create a one-slide deck for an alternate table content structure."""
    return Presentation(
        id=identifier,
        title="Table Structure Test",
        metadata=PresentationMetadata(root={}),
        slides=[Slide(
            id="comparison-slide",
            order=1,
            layout=SlideLayout.COMPARISON,
            purpose="Check table header and row normalization",
            components=[
                component("comparison-title", ComponentType.TITLE, "Comparison"),
                component("comparison-table", ComponentType.TABLE, table_content),
            ],
        )],
    )


def main() -> None:
    output = Path(__file__).parent / "generated" / "test_rendered_presentation.pptx"
    sample = sample_presentation()
    fixture_table = next(component for component in sample.slides[3].components if component.type == ComponentType.TABLE)
    print("Deterministic headers + flat rows table component:")
    print(fixture_table.model_dump_json(indent=2))
    path = render_presentation(sample, output)
    assert path.is_file(), f"Renderer did not create output file: {path}"
    assert path.stat().st_size > 0, "Rendered PowerPoint is empty"
    reopened = PowerPoint(path)
    assert len(reopened.slides) == 9, f"Expected 9 slides, found {len(reopened.slides)}"

    card_text = "\n".join(shape.text for shape in reopened.slides[4].shapes if shape.has_text_frame)
    for expected in ("Solar", "Wind", "Hydro", "Harnessing sunlight", "moving air", "flowing water"):
        assert expected in card_text, f"Three-card slide is missing: {expected}"
    assert "Point 1" not in card_text, "Structured card heading was replaced with a generic label"

    comparison_table = next(shape.table for shape in reopened.slides[3].shapes if shape.has_table)
    table_text = [cell.text for row in comparison_table.rows for cell in row.cells]
    for expected in ("Source", "Primary Driver", "Availability", "Solar", "Sunlight", "Daytime only", "Wind", "Air currents", "Variable", "Hydro", "Water flow", "Consistent", "Geothermal", "Heat from the Earth", "Location-dependent", "Biomass", "Organic material", "Fuel availability", "Tidal", "Ocean tides", "Coastal variation", "Wave", "Ocean surface motion", "Site-dependent"):
        assert expected in table_text, f"Comparison table is missing: {expected}"
    assert len(comparison_table.rows) == 8, f"Expected header plus all 7 grouped rows, found {len(comparison_table.rows)} rows"

    columns_output = Path(__file__).parent / "generated" / "_test_columns_table.pptx"
    encoded_output = Path(__file__).parent / "generated" / "_test_encoded_table.pptx"
    try:
        render_presentation(table_presentation("columns_table_test", {
            "columns": ["Technology", "Energy Input", "Availability"],
            "rows": [
                "Solar", "Sunlight", "Daytime",
                "Wind", "Air currents", "Variable",
            ],
        }), columns_output)
        columns_table = next(shape.table for shape in PowerPoint(columns_output).slides[0].shapes if shape.has_table)
        columns_text = [cell.text for row in columns_table.rows for cell in row.cells]
        for expected in ("Technology", "Energy Input", "Availability", "Solar", "Sunlight", "Daytime", "Wind", "Air currents", "Variable"):
            assert expected in columns_text, f"Columns + flat rows table is missing: {expected}"
        assert len(columns_table.rows) == 3, "Columns + flat rows did not produce a header and both data rows"

        render_presentation(table_presentation("encoded_table_test", {
            "content": json.dumps({
                "headers": ["Factor", "Result"],
                "rows": ["Input", "Output", "Cause", "Effect"],
            }),
        }), encoded_output)
        encoded_table = next(shape.table for shape in PowerPoint(encoded_output).slides[0].shapes if shape.has_table)
        encoded_text = [cell.text for row in encoded_table.rows for cell in row.cells]
        for expected in ("Factor", "Result", "Input", "Output", "Cause", "Effect"):
            assert expected in encoded_text, f"Nested JSON table is missing: {expected}"
    finally:
        columns_output.unlink(missing_ok=True)
        encoded_output.unlink(missing_ok=True)

    statistic_slide = reopened.slides[7]
    statistic_label = next(shape for shape in statistic_slide.shapes if shape.has_text_frame and "deliberately long statistic label" in shape.text)
    assert statistic_label.text_frame.word_wrap, "Long statistic label does not wrap"
    assert statistic_label.text_frame.auto_size == MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE, "Long statistic label is not set to fit its area"
    assert statistic_label.height <= Inches(2), "Statistic label exceeds its bounded slide area"

    conclusion_slide = reopened.slides[8]
    conclusion_text = "\n".join(shape.text for shape in conclusion_slide.shapes if shape.has_text_frame)
    assert "Renewable energy can support a more diverse electricity system." in conclusion_text
    assert "Use resources suited to each region" in conclusion_text
    conclusion_summary = next(shape for shape in conclusion_slide.shapes if shape.has_text_frame and "Renewable energy can support" in shape.text)
    assert conclusion_summary.height >= Inches(1.5), "Conclusion text does not have a readable content area"
    assert conclusion_summary.top + conclusion_summary.height <= Inches(6.5), "Conclusion text extends beyond the slide content area"

    print(f"python-pptx: {version('python-pptx')}")
    print(f"Rendered: {path}")
    print(f"File size: {path.stat().st_size:,} bytes")
    print(f"Slides reopened: {len(reopened.slides)}")
    print("Alternate columns + flat rows and nested JSON table structures reopened successfully.")
    print("Layouts verified: " + ", ".join(layout.value for layout in SlideLayout))


if __name__ == "__main__":
    main()
