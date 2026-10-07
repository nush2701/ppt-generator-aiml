"""Live Gemini checks for slide-level selective regeneration."""

from pathlib import Path
from uuid import uuid4

from pptx import Presentation as PowerPoint

from app.models.presentation import (
    ComponentType,
    Presentation,
    PresentationMetadata,
    Slide,
    SlideComponent,
    SlideLayout,
)
from app.models.requirements import PresentationRequirements
from app.models.selective_regeneration import SelectiveRegenerationStatus
from app.services.selective_regenerator import SelectiveRegenerator


def make_presentation() -> Presentation:
    """Build a stable, identifiable five-slide source deck."""
    titles = [
        "Introduction",
        "Solar Energy",
        "Renewable Energy Comparison",
        "Energy Transition",
        "Conclusion",
    ]
    layouts = [
        SlideLayout.TITLE,
        SlideLayout.TITLE_BULLETS,
        SlideLayout.COMPARISON,
        SlideLayout.TITLE_BULLETS,
        SlideLayout.CONCLUSION,
    ]
    slides = []
    for index, (title, layout) in enumerate(zip(titles, layouts, strict=True), start=1):
        components = [
            SlideComponent(
                id=f"slide-{index}-title",
                type=ComponentType.TITLE,
                content=title,
            )
        ]
        if index == 3:
            components.append(
                SlideComponent(
                    id="slide-3-table",
                    type=ComponentType.TABLE,
                    content={
                        "headers": ["Source", "Driver"],
                        "rows": ["Solar", "Sunlight", "Wind", "Air currents"],
                    },
                )
            )
        elif index > 1:
            components.append(
                SlideComponent(
                    id=f"slide-{index}-body",
                    type=ComponentType.TEXT,
                    content=f"Stable existing content for slide {index}.",
                )
            )
        slides.append(
            Slide(
                id=f"slide-{index}",
                order=index,
                layout=layout,
                purpose=f"Explain {title.lower()}.",
                components=components,
            )
        )

    return Presentation(
        id=f"presentation-selective-test-{uuid4().hex}",
        title="Renewable Energy",
        metadata=PresentationMetadata(
            root={
                "topic": "renewable energy",
                "purpose": None,
                "audience": "college students",
                "slide_count": 5,
                "tone": "professional",
                "language": "English",
                "key_requirements": [],
                "constraints": [],
            }
        ),
        slides=slides,
    )


REQUIREMENTS = PresentationRequirements(
    topic="renewable energy",
    purpose=None,
    audience="college students",
    slide_count=5,
    tone="professional",
    language="English",
    key_requirements=[],
    constraints=[],
)


def serialized_slides(presentation: Presentation) -> dict[str, dict[str, object]]:
    return {
        slide.id: slide.model_dump(mode="json")
        for slide in presentation.slides
    }


def main() -> None:
    presentation = make_presentation()
    before = serialized_slides(presentation)
    presentation_id = presentation.id
    output_path = Path(__file__).parent / "generated" / f"selective_regeneration_{uuid4().hex}.pptx"

    service = SelectiveRegenerator()
    result = service.selectively_regenerate(
        presentation=presentation,
        requirements=REQUIREMENTS,
        modification_request="Change slide 3 to compare solar and wind energy instead.",
        output_path=output_path,
    )
    assert result.status == SelectiveRegenerationStatus.REGENERATED, result.message
    assert result.presentation.id == presentation_id
    assert result.affected_slides == ["slide-3"], result.affected_slides
    assert result.regenerated_slides == ["slide-3"], result.regenerated_slides
    assert result.preserved_slides == ["slide-1", "slide-2", "slide-4", "slide-5"]
    regenerated_slide = next(slide for slide in result.presentation.slides if slide.id == "slide-3")
    assert regenerated_slide.order == 3
    assert regenerated_slide.layout == SlideLayout.COMPARISON

    after = serialized_slides(result.presentation)
    for slide_id in result.preserved_slides:
        assert before[slide_id] == after[slide_id], f"Unaffected slide changed: {slide_id}"
    assert serialized_slides(presentation) == before, "The original presentation was mutated."

    assert output_path.is_file(), f"PowerPoint was not created: {output_path}"
    assert output_path.stat().st_size > 0, "PowerPoint file is empty"
    reopened = PowerPoint(output_path)
    assert len(reopened.slides) == 5, f"Expected 5 slides, found {len(reopened.slides)}"

    unclear_source = make_presentation()
    unclear_before = serialized_slides(unclear_source)
    unclear = service.selectively_regenerate(
        unclear_source,
        REQUIREMENTS,
        "Change it.",
    )
    assert unclear.status == SelectiveRegenerationStatus.CLARIFICATION_NEEDED, unclear.message
    assert unclear.clarification_needed
    assert not unclear.regenerated_slides
    assert serialized_slides(unclear.presentation) == unclear_before
    assert serialized_slides(unclear_source) == unclear_before
    assert unclear.output_path is None

    global_source = make_presentation()
    global_before = serialized_slides(global_source)
    global_result = service.selectively_regenerate(
        global_source,
        REQUIREMENTS,
        "Make the entire presentation more professional.",
    )
    assert global_result.status == SelectiveRegenerationStatus.GLOBAL_CHANGE, global_result.message
    assert not global_result.regenerated_slides
    assert global_result.preserved_slides == [f"slide-{number}" for number in range(1, 6)]
    assert serialized_slides(global_result.presentation) == global_before
    assert serialized_slides(global_source) == global_before
    assert global_result.output_path is None

    print(f"Selective regeneration status: {result.status.value}")
    print(f"Affected slides: {result.affected_slides}")
    print(f"Regenerated slides: {result.regenerated_slides}")
    print(f"Preserved slides: {result.preserved_slides}")
    print(f"Presentation ID preserved: {result.presentation.id == presentation_id}")
    print(f"Clarification case: {unclear.status.value}")
    print(f"Global case: {global_result.status.value}")
    print(f"PowerPoint path: {output_path}")
    print(f"PowerPoint size: {output_path.stat().st_size:,} bytes")
    print("Live Gemini selective-regeneration checks passed.")


if __name__ == "__main__":
    main()
